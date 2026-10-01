"""Administración, configuración de sede, auditoría y reportes."""

import datetime as dt
import io
from collections.abc import Callable

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.clock import FixedClock

from .conftest import Actor, api, reason_id, register, site_id


def test_user_lifecycle_with_temporary_password(
    client: TestClient, make_user: Callable[..., Actor], db: Session
) -> None:
    admin = make_user("ADMIN", sites=("ALZ", "BAR"))
    created = client.post(
        api("/admin/users"),
        json={
            "username": "Maria.Lopez",
            "full_name": "María López",
            "role_codes": ["OPERATOR"],
            "site_ids": [site_id(db, "BAR")],
        },
        headers=admin.headers,
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert (body["username"], body["must_change_password"], body["roles"]) == ("maria.lopez", True, ["OPERATOR"])
    temporary = body["temporary_password"]

    login = client.post(api("/auth/login"), json={"username": "maria.lopez", "password": temporary})
    assert login.json()["user"]["must_change_password"] is True

    duplicate = client.post(
        api("/admin/users"),
        json={"username": "maria.lopez", "full_name": "Otra", "role_codes": ["OPERATOR"]},
        headers=admin.headers,
    )
    assert duplicate.json()["code"] == "USERNAME_TAKEN"

    deactivated = client.patch(
        api(f"/admin/users/{body['public_id']}"), json={"is_active": False}, headers=admin.headers
    )
    assert deactivated.json()["is_active"] is False
    token = login.json()["access_token"]
    assert client.get(api("/auth/me"), headers={"Authorization": f"Bearer {token}"}).status_code == 401

    history = client.get(api("/audit-events"), params={"resource_id": body["public_id"]}, headers=admin.headers).json()
    assert {e["action"] for e in history["items"]} >= {"USER_CREATE", "USER_UPDATE"}


def test_admin_cannot_lock_itself_out(client: TestClient, make_user: Callable[..., Actor], db: Session) -> None:
    admin = make_user("ADMIN")
    me = client.patch(api(f"/admin/users/{admin.public_id}"), json={"is_active": False}, headers=admin.headers)
    assert me.status_code == 422
    roles = client.get(api("/admin/roles"), headers=admin.headers).json()
    admin_role = next(r for r in roles if r["code"] == "ADMIN")
    stripped = client.put(
        api(f"/admin/roles/{admin_role['id']}/permissions"),
        json={"permission_codes": ["user:read"]},
        headers=admin.headers,
    )
    assert stripped.status_code == 422


def test_parameters_are_typed_and_audited(client: TestClient, make_user: Callable[..., Actor]) -> None:
    admin = make_user("ADMIN")
    bad = client.put(api("/admin/parameters/auth.lockout_minutes"), json={"value": "quince"}, headers=admin.headers)
    assert bad.json()["code"] == "PARAMETER_TYPE_MISMATCH"
    locked = client.put(api("/admin/parameters/app.timezone"), json={"value": "UTC"}, headers=admin.headers)
    assert locked.json()["code"] == "PARAMETER_NOT_EDITABLE"
    ok = client.put(api("/admin/parameters/auth.lockout_minutes"), json={"value": 20}, headers=admin.headers)
    assert ok.json()["value"] == 20
    client.put(api("/admin/parameters/auth.lockout_minutes"), json={"value": 15}, headers=admin.headers)


def test_templates_are_validated_before_saving(client: TestClient, make_user: Callable[..., Actor]) -> None:
    admin = make_user("ADMIN")
    templates = client.get(api("/admin/notification-templates"), headers=admin.headers).json()
    called = next(t for t in templates if t["code"] == "APPT_CALLED")
    broken = client.patch(
        api(f"/admin/notification-templates/{called['id']}"),
        json={"body_text": "Hola {{ worker_first_name "},
        headers=admin.headers,
    )
    assert broken.json()["code"] == "TEMPLATE_INVALID"
    preview = client.get(api(f"/admin/notification-templates/{called['id']}/preview"), headers=admin.headers)
    assert "A-007" in preview.json()["subject"]


def test_site_settings_versioning(
    client: TestClient, make_user: Callable[..., Actor], make_worker: Callable[..., str], db: Session, clock: FixedClock
) -> None:
    supervisor = make_user("SUPERVISOR")
    alz = site_id(db)
    payload = {"daily_capacity": 25, "slot_minutes": 20, "tolerance_minutes": 5, "change_reason": "Nuevo médico"}
    today = client.post(
        api(f"/sites/{alz}/settings"),
        json={**payload, "valid_from": clock.today().isoformat()},
        headers=supervisor.headers,
    )
    assert today.json()["code"] == "DATE_NOT_ALLOWED"  # RN-08: aplica desde mañana
    far = (clock.today() + dt.timedelta(days=400)).isoformat()
    created = client.post(
        api(f"/sites/{alz}/settings"), json={**payload, "valid_from": far}, headers=supervisor.headers
    )
    assert created.status_code == 201
    settings = client.get(api(f"/sites/{alz}/settings"), headers=supervisor.headers).json()
    assert settings["current"]["daily_capacity"] == 20  # hoy sigue la configuración vigente
    assert any(h["valid_from"] == far for h in settings["history"])


def test_schedule_replacement(
    client: TestClient, make_user: Callable[..., Actor], db: Session, clock: FixedClock
) -> None:
    supervisor = make_user("SUPERVISOR", sites=("BAR",))
    bar = site_id(db, "BAR")
    valid_from = (clock.today() + dt.timedelta(days=3000)).isoformat()
    blocks = [{"weekday": d, "block": "AM", "start_time": "08:30", "end_time": "13:00"} for d in range(1, 6)]
    response = client.put(
        api(f"/sites/{bar}/schedules"),
        headers=supervisor.headers,
        json={"valid_from": valid_from, "blocks": blocks, "change_reason": "Horario de verano"},
    )
    assert response.status_code == 200, response.text
    assert len(response.json()) == 5
    overlap = [
        {"weekday": 1, "block": "AM", "start_time": "08:00", "end_time": "12:00"},
        {"weekday": 1, "block": "PM", "start_time": "11:00", "end_time": "14:00"},
    ]
    later = (clock.today() + dt.timedelta(days=3100)).isoformat()
    bad = client.put(
        api(f"/sites/{bar}/schedules"),
        headers=supervisor.headers,
        json={"valid_from": later, "blocks": overlap, "change_reason": "Horario con cruce"},
    )
    assert bad.json()["code"] == "SCHEDULE_OVERLAP"


def test_audit_chain_verification_endpoint(client: TestClient, make_user: Callable[..., Actor]) -> None:
    auditor = make_user("AUDITOR")
    result = client.post(api("/audit-events/verify"), headers=auditor.headers).json()
    assert result["intact"] is True
    assert result["events_checked"] > 0


def test_reports_and_exports(
    client: TestClient,
    make_user: Callable[..., Actor],
    operator: Actor,
    make_worker: Callable[..., str],
    db: Session,
    clock: FixedClock,
) -> None:
    alz = site_id(db)
    a = register(client, operator, alz, make_worker()).json()
    b = register(client, operator, alz, make_worker()).json()
    client.post(api(f"/appointments/{a['public_id']}/call"), json={}, headers=operator.headers)
    clock.advance(minutes=5)
    client.post(api(f"/appointments/{a['public_id']}/start"), json={}, headers=operator.headers)
    clock.advance(minutes=8)  # < 15 min: el access token sigue vigente
    client.post(api(f"/appointments/{a['public_id']}/finish"), json={}, headers=operator.headers)
    client.post(
        api(f"/appointments/{b['public_id']}/cancel"),
        json={"reason_id": reason_id(client, operator, "CANCEL", "YA_NO_REQUIERE")},
        headers=operator.headers,
    )

    day = clock.today().isoformat()
    supervisor = make_user("SUPERVISOR")
    summary = client.get(
        api("/reports/summary"), params={"date_from": day, "date_to": day, "site_id": alz}, headers=supervisor.headers
    ).json()
    totals = summary["totals"]
    assert (totals["requested"], totals["attended"], totals["cancelled"], totals["capacity"]) == (2, 1, 1, 20)
    assert totals["avg_service_minutes"] == 8.0
    assert totals["avg_wait_minutes"] == 5.0
    assert summary["by_day"][0]["attended"] == 1

    csv_export = client.get(
        api("/reports/export"), params={"format": "csv", "date_from": day, "date_to": day}, headers=supervisor.headers
    )
    assert csv_export.status_code == 200
    assert "attended" in csv_export.text
    nominal = client.get(
        api("/reports/export"),
        params={"report": "appointments", "format": "xlsx", "date_from": day, "date_to": day},
        headers=supervisor.headers,
    )
    assert nominal.headers["content-type"].startswith("application/vnd.openxmlformats")

    # Reporte de indicadores completo: Excel con varias hojas y PDF institucional
    xlsx = client.get(api("/reports/export"), params={"date_from": day, "date_to": day}, headers=supervisor.headers)
    sheets = load_workbook(io.BytesIO(xlsx.content)).sheetnames
    assert sheets == ["Resumen", "Por día", "Por médico", "Por dependencia", "Por canal", "Por hora"]
    for report in ("summary", "appointments"):
        pdf = client.get(
            api("/reports/export"),
            params={"report": report, "format": "pdf", "date_from": day, "date_to": day},
            headers=supervisor.headers,
        )
        assert (pdf.status_code, pdf.headers["content-type"]) == (200, "application/pdf")
        assert pdf.content.startswith(b"%PDF")

    # La encargada descarga el reporte agregado (sin datos personales), pero no el listado nominal
    assert client.get(api("/reports/summary"), headers=operator.headers).status_code == 200
    assert client.get(api("/reports/export"), params={"format": "pdf"}, headers=operator.headers).status_code == 200
    denied = client.get(api("/reports/export"), params={"report": "appointments"}, headers=operator.headers)
    assert denied.status_code == 403
