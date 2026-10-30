"""Reporte de satisfacción (rating:read): solo jefatura; filtros, paginación y exportación con comentarios."""

import io
import re
from collections.abc import Callable

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modules.notifications.models import Notification

from .conftest import Actor, api, register, site_id

WEEK = [{"weekday": d, "block": "AM", "start_time": "08:00", "end_time": "13:00"} for d in range(1, 6)]


@pytest.fixture
def site(client: TestClient, make_user: Callable[..., Actor], request: pytest.FixtureRequest) -> str:
    code = "SA" + "".join(chr(65 + int(c)) for c in f"{abs(hash(request.node.name)) % 10**6:06d}")
    admin = make_user("ADMIN", sites=("ALZ",))
    body = {"code": code, "name": f"Sede {code}", "short_name": code, "ticket_prefix": code[-3:], "blocks": WEEK}
    assert client.post(api("/sites"), json=body, headers=admin.headers).status_code == 201
    return code


def rate(
    client: TestClient, app: FastAPI, db: Session, operator: Actor, sid: int, dni: str, answer: dict[str, object]
) -> None:
    """Atiende a un trabajador y responde la encuesta anónima que le llega por correo."""
    original = app.state.settings
    app.state.settings = original.model_copy(update={"public_app_url": "http://topico.local/consulta"})
    try:
        appt = register(client, operator, sid, dni).json()
        client.post(api(f"/sites/{sid}/queue/call-next"), headers=operator.headers)
        client.post(api(f"/appointments/{appt['public_id']}/start"), json={}, headers=operator.headers)
        client.post(api(f"/appointments/{appt['public_id']}/finish"), json={}, headers=operator.headers)
    finally:
        app.state.settings = original
    mail = db.scalar(
        select(Notification).where(Notification.template_code == "RATING_REQUEST").order_by(Notification.id.desc())
    )
    assert mail is not None
    assert mail.body_text
    match = re.search(r"/calificar#([\w-]+)", mail.body_text)
    assert match
    sent = client.post(api("/public/rating"), json={"token": match.group(1), **answer})
    assert sent.status_code == 204, sent.text


def test_satisfaction_report_for_supervision_only(
    client: TestClient,
    app: FastAPI,
    make_user: Callable[..., Actor],
    make_worker: Callable[..., str],
    db: Session,
    site: str,
) -> None:
    sid = site_id(db, site)
    operator = make_user("OPERATOR", sites=(site,))
    answers = [
        {"score": 5, "wait_score": 4, "comment": "Muy amables"},
        {"score": 2, "wait_score": 1, "comment": "Esperé mucho <b>& sin aviso</b>"},
        {"score": 4, "wait_score": None, "comment": None},
    ]
    for answer in answers:
        rate(client, app, db, operator, sid, make_worker(), answer)

    url = api("/reports/satisfaction")
    # La encargada no ve el reporte ni los comentarios (en Reportes ve solo los promedios)
    assert client.get(url, params={"site_id": sid}, headers=operator.headers).status_code == 403
    summary = client.get(api("/reports/summary"), params={"site_id": sid}, headers=operator.headers).json()["rating"]
    assert (summary["count"], summary["comments"]) == (3, [])

    supervisor = make_user("SUPERVISOR", sites=(site,))
    body = client.get(url, params={"site_id": sid}, headers=supervisor.headers).json()
    assert (body["invited"], body["count"], body["with_comment"], body["response_pct"]) == (3, 3, 2, 100.0)
    assert (body["avg_score"], body["avg_wait_score"]) == (3.67, 2.5)
    assert [b["count"] for b in body["distribution"]] == [0, 1, 0, 1, 1]
    assert body["by_site"][0] | {"site_name": None} == {
        "site_name": None,
        "count": 3,
        "avg_score": 3.67,
        "avg_wait_score": 2.5,
    }
    assert len(body["by_month"]) == 1
    assert body["comments_total"] == 2
    assert {c["comment"] for c in body["comments"]} == {"Muy amables", "Esperé mucho <b>& sin aviso</b>"}
    assert "dni" not in str(body).lower()

    low = client.get(url, params={"site_id": sid, "score": 2}, headers=supervisor.headers).json()
    assert [c["score"] for c in low["comments"]] == [2]
    paged = client.get(url, params={"site_id": sid, "page_size": 1, "page": 2}, headers=supervisor.headers).json()
    assert (paged["comments_total"], len(paged["comments"])) == (2, 1)

    xlsx = client.get(f"{url}/export", params={"site_id": sid, "format": "xlsx"}, headers=supervisor.headers)
    assert xlsx.status_code == 200, xlsx.text
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Resumen", "Por mes", "Por sede", "Por médico", "Comentarios"]
    texts = [row[5] for row in book["Comentarios"].iter_rows(values_only=True)]
    assert "Muy amables" in texts
    pdf = client.get(f"{url}/export", params={"site_id": sid, "format": "pdf"}, headers=supervisor.headers)
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")  # el comentario con «<b>&» no rompe el PDF

    # Sede no autorizada
    other = make_user("SUPERVISOR", sites=("ALZ",))
    assert client.get(url, params={"site_id": sid}, headers=other.headers).status_code == 403


def test_me_reports_idle_timeout(client: TestClient, make_user: Callable[..., Actor]) -> None:
    user = make_user("OPERATOR")
    assert client.get(api("/auth/me"), headers=user.headers).json()["idle_timeout_minutes"] == 20
