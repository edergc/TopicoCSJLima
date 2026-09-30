"""Lectura y normalización del Excel de trabajadores EPS Rímac.

Los encabezados se reconocen por sinónimos (sin tildes, mayúsculas/minúsculas indistintas),
porque los Excel institucionales varían. Nunca se escribe en tablas finales desde aquí.
"""

import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.core.errors import BusinessRuleError
from app.shared.text import clean_name, is_valid_dni, normalize_document, normalize_key

HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "document_number": (
        "dni",
        "n_dni",
        "nro_dni",
        "numero_dni",
        "num_dni",
        "documento",
        "nro_documento",
        "numero_documento",
        "doc_identidad",
        "documento_identidad",
        "n_documento",
    ),
    "first_names": ("nombres", "nombre", "primer_nombre"),
    "paternal_surname": ("apellido_paterno", "ap_paterno", "apellidopaterno", "paterno", "apell_paterno"),
    "maternal_surname": ("apellido_materno", "ap_materno", "apellidomaterno", "materno", "apell_materno"),
    "surnames": ("apellidos",),
    "full_name": (
        "apellidos_y_nombres",
        "apellidos_nombres",
        "nombres_y_apellidos",
        "nombre_completo",
        "trabajador",
        "colaborador",
    ),
    "sex": ("sexo", "genero"),
    "birth_date": ("fecha_nacimiento", "fecha_de_nacimiento", "fec_nac", "f_nacimiento", "fnacimiento", "f_nac"),
    "institutional_email": ("correo", "correo_institucional", "email", "e_mail", "correo_electronico", "mail"),
    "phone": ("telefono", "celular", "movil", "nro_telefono", "telefono_celular", "nro_celular"),
    "department_name": (
        "dependencia",
        "organo",
        "organo_jurisdiccional",
        "area",
        "unidad",
        "unidad_organica",
        "oficina",
    ),
    "employee_code": ("codigo", "cod_trabajador", "codigo_trabajador", "cod_planilla", "codigo_planilla"),
    "age": ("edad",),
}
_SYNONYM_INDEX = {syn: target for target, syns in HEADER_SYNONYMS.items() for syn in syns}
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_PHONE_CLEAN = re.compile(r"[^\d+]")
MAX_ROWS = 20_000

WORKER_FIELDS = (
    "first_names",
    "paternal_surname",
    "maternal_surname",
    "sex",
    "birth_date",
    "institutional_email",
    "phone",
    "department_name",
    "employee_code",
)


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", normalize_key(str(value or "")).lower()).strip("_")


@dataclass
class ParsedRow:
    row_number: int
    raw: dict[str, Any]
    normalized: dict[str, Any] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class ParsedSheet:
    sheet_name: str
    mapping: dict[str, str]  # encabezado original → campo
    rows: list[ParsedRow]
    ignored_columns: list[str]


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _jsonable(value: object) -> Any:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def read_workbook(content: bytes) -> ParsedSheet:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise BusinessRuleError("IMPORT_FILE_INVALID") from exc
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            for index, header in enumerate(rows, start=1):
                if index > 10:
                    break
                mapping = {
                    str(h): _SYNONYM_INDEX[header_key(h)] for h in header if h and header_key(h) in _SYNONYM_INDEX
                }
                if "document_number" in mapping.values() and len(mapping) >= 2:
                    return _parse_rows(sheet.title, list(header), mapping, rows, index)
        raise BusinessRuleError(
            "IMPORT_COLUMNS_MISSING",
            "No se encontró una fila de encabezados con la columna DNI y los nombres del trabajador.",
        )
    finally:
        workbook.close()


def _parse_rows(title: str, header: list[object], mapping: dict[str, str], rows: Any, header_row: int) -> ParsedSheet:
    fields = set(mapping.values())
    has_names = (
        {"first_names", "paternal_surname"} <= fields
        or "full_name" in fields
        or ("first_names" in fields and "surnames" in fields)
    )
    if not has_names:
        raise BusinessRuleError(
            "IMPORT_COLUMNS_MISSING",
            "Faltan columnas de nombres: se requiere NOMBRES y APELLIDO PATERNO (o APELLIDOS Y NOMBRES).",
        )
    ignored = [str(h) for h in header if h and str(h) not in mapping]
    parsed: list[ParsedRow] = []
    for offset, values in enumerate(rows, start=1):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        if len(parsed) >= MAX_ROWS:
            raise BusinessRuleError("IMPORT_FILE_INVALID", f"El archivo supera el máximo de {MAX_ROWS} filas.")
        raw = {str(h): _jsonable(v) for h, v in zip(header, values, strict=False) if h is not None}
        row = ParsedRow(row_number=header_row + offset, raw=raw)
        by_field = {
            mapping[str(h)]: v for h, v in zip(header, values, strict=False) if h is not None and str(h) in mapping
        }
        normalize_row(row, by_field)
        parsed.append(row)
    if not parsed:
        raise BusinessRuleError("IMPORT_EMPTY")
    return ParsedSheet(title, mapping, parsed, ignored)


def _parse_date(value: object) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def normalize_row(row: ParsedRow, values: dict[str, object]) -> None:
    n: dict[str, Any] = {}

    # Documento
    raw_doc = values.get("document_number")
    dni = normalize_document(raw_doc)
    if not dni:
        row.errors.append("DNI vacío.")
    elif not is_valid_dni(dni):
        row.errors.append(f"DNI inválido: '{_cell_text(raw_doc)}' (debe tener 8 dígitos).")
    else:
        original = _cell_text(raw_doc) or ""
        if original.replace(" ", "") != dni and original.isdigit():
            row.warnings.append(f"DNI completado con ceros iniciales: {original} → {dni}.")
    n["document_number"] = dni

    # Nombres
    first = clean_name(_cell_text(values.get("first_names")))
    paternal = clean_name(_cell_text(values.get("paternal_surname")))
    maternal = clean_name(_cell_text(values.get("maternal_surname")))
    if not paternal and values.get("surnames"):
        parts = (clean_name(_cell_text(values.get("surnames"))) or "").split(" ", 1)
        paternal, maternal = parts[0] or None, (parts[1] if len(parts) > 1 else None)
    if (not first or not paternal) and values.get("full_name"):
        full = clean_name(_cell_text(values.get("full_name"))) or ""
        if "," in full:
            surnames, first = (p.strip() for p in full.split(",", 1))
            parts = surnames.split(" ", 1)
            paternal, maternal = parts[0], (parts[1] if len(parts) > 1 else None)
        else:
            row.errors.append("Nombre completo sin separar: use el formato 'APELLIDOS, NOMBRES'.")
    if not first:
        row.errors.append("Nombres vacíos.")
    if not paternal:
        row.errors.append("Apellido paterno vacío.")
    n["first_names"], n["paternal_surname"], n["maternal_surname"] = first, paternal, maternal

    # Sexo
    sex_raw = normalize_key(_cell_text(values.get("sex")) or "")
    sex = {"F": "F", "FEMENINO": "F", "MUJER": "F", "M": "M", "MASCULINO": "M", "HOMBRE": "M"}.get(sex_raw)
    if sex_raw and sex is None:
        row.warnings.append(f"Sexo no reconocido: '{sex_raw}' (se deja vacío).")
    n["sex"] = sex

    # Fecha de nacimiento
    try:
        birth = _parse_date(values.get("birth_date"))
        if birth and (birth.year < 1900 or birth > date.today()):
            row.warnings.append("Fecha de nacimiento fuera de rango (se deja vacía).")
            birth = None
        n["birth_date"] = birth.isoformat() if birth else None
    except ValueError:
        row.warnings.append(f"Fecha de nacimiento no reconocida: '{values.get('birth_date')}' (se deja vacía).")
        n["birth_date"] = None

    # Contacto
    email = (_cell_text(values.get("institutional_email")) or "").lower() or None
    if email and not _EMAIL_RE.match(email):
        row.warnings.append(f"Correo inválido: '{email}' (se deja vacío).")
        email = None
    n["institutional_email"] = email
    phone = _PHONE_CLEAN.sub("", _cell_text(values.get("phone")) or "") or None
    if phone and not (6 <= len(phone) <= 20):
        row.warnings.append(f"Teléfono inválido: '{phone}' (se deja vacío).")
        phone = None
    n["phone"] = phone

    department = _cell_text(values.get("department_name"))
    n["department_name"] = " ".join(department.split()).upper() if department else None
    code = _cell_text(values.get("employee_code"))
    n["employee_code"] = code[:20] if code else None
    if values.get("age") is not None:
        row.warnings.append("La columna EDAD se ignora (la edad se calcula desde la fecha de nacimiento).")
    row.normalized = n
