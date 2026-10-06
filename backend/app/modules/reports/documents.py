"""Documentos descargables de reportes: Excel con formato y PDF institucional.

El reporte de indicadores (resumen) no contiene datos personales. El listado nominal sí:
solo se genera con el permiso report:read_nominal (lo valida el router).
"""

import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.core.clock import INSTITUTION_TZ

BRAND = colors.HexColor("#7A1E2C")
BRAND_DARK = colors.HexColor("#45111A")
BRAND_SOFT = colors.HexColor("#FBF4F5")
INK = colors.HexColor("#1C1917")
MUTED = colors.HexColor("#57534E")
LINE = colors.HexColor("#E7E5E4")
CHANNEL_LABEL = {"PHONE": "Teléfono", "WALK_IN": "Presencial", "MIGRATION": "Migración"}
STATUS_LABEL = {
    "REGISTRADO": "Registrado",
    "EN_ESPERA": "En espera",
    "LLAMADO": "Llamado",
    "EN_ATENCION": "En atención",
    "ATENDIDO": "Atendido",
    "CANCELADO": "Cancelado",
    "NO_PRESENTADO": "No presentado",
    "ANULADO": "Anulado",
}


@dataclass(frozen=True)
class ReportMeta:
    institution: str
    date_from: date
    date_to: date
    site_names: list[str]
    generated_at: datetime
    generated_by: str

    @property
    def period(self) -> str:
        return f"{self.date_from:%d/%m/%Y} al {self.date_to:%d/%m/%Y}"

    @property
    def sites(self) -> str:
        return ", ".join(self.site_names) or "—"

    @property
    def generated(self) -> str:
        return f"{self.generated_at.astimezone(INSTITUTION_TZ):%d/%m/%Y %H:%M} por {self.generated_by}"


def _num(value: float | int | None, suffix: str = "") -> str:
    if value is None:
        return "—"
    text = f"{value:,.1f}" if isinstance(value, float) else f"{value:,}"
    return text.replace(",", " ") + suffix


def _kpis(totals: dict[str, Any]) -> list[tuple[str, str]]:
    return [
        ("Capacidad ofertada", _num(totals["capacity"])),
        ("Solicitudes", _num(totals["requested"])),
        ("Atendidos", _num(totals["attended"])),
        ("No presentados", _num(totals["no_show"])),
        ("Cancelados", _num(totals["cancelled"])),
        ("Utilización", _num(totals["utilization_pct"], " %")),
        ("Espera promedio", _num(totals["avg_wait_minutes"], " min")),
        ("Atención promedio", _num(totals["avg_service_minutes"], " min")),
    ]


def _rating_rows(rating: dict[str, Any]) -> list[list[Any]]:
    """Satisfacción (anónima): agregados y distribución de puntajes."""
    rows: list[list[Any]] = [
        ["Calificaciones recibidas", f"{rating['count']} de {rating['invited']} invitaciones"],
        ["Trato (promedio de 1 a 5)", _num(rating["avg_score"])],
        ["Tiempo de espera (promedio de 1 a 5)", _num(rating["avg_wait_score"])],
    ]
    rows += [[f"Calificaciones con {b['score']} estrella(s)", b["count"]] for b in rating["distribution"]]
    return rows


def _day_rows(summary: dict[str, Any]) -> list[list[Any]]:
    return [
        [d["service_date"], d["site_name"], d["capacity"], d["requested"], d["attended"], d["no_show"], d["cancelled"]]
        for d in summary["by_day"]
    ]


DAY_HEADERS = ["Fecha", "Sede", "Capacidad", "Solicitudes", "Atendidos", "No present.", "Cancelados"]


# =============================================================================== Excel
_XL_HEADER_FILL = PatternFill("solid", fgColor="7A1E2C")
_XL_HEADER_FONT = Font(bold=True, color="FFFFFF")
_XL_THIN = Side(style="thin", color="E7E5E4")


def _xl_table(
    ws: Worksheet, row: int, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None
) -> int:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.fill, cell.font = _XL_HEADER_FILL, _XL_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, values in enumerate(rows, row + 1):
        for c, value in enumerate(values, 1):
            if isinstance(value, datetime):
                value = value.astimezone(INSTITUTION_TZ).replace(tzinfo=None)
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = Border(bottom=_XL_THIN)
            if isinstance(value, date) and not isinstance(value, datetime):
                cell.number_format = "DD/MM/YYYY"
            elif isinstance(value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
    for col, width in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    return row + len(rows) + 1


def _xl_title(ws: Worksheet, title: str, meta: ReportMeta) -> int:
    ws["A1"] = meta.institution.upper()
    ws["A1"].font = Font(bold=True, size=10, color="7A1E2C")
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=14)
    ws["A3"] = f"Periodo: {meta.period}   ·   Sede(s): {meta.sites}"
    ws["A4"] = f"Generado el {meta.generated}"
    ws["A4"].font = Font(italic=True, size=9, color="78716C")
    return 6


def summary_xlsx(summary: dict[str, Any], meta: ReportMeta) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Resumen"
    row = _xl_title(ws, "Reporte de atenciones del Tópico de Salud", meta)
    row = _xl_table(ws, row, ["Indicador", "Valor"], [list(k) for k in _kpis(summary["totals"])], [34, 18])
    ws.cell(row=row + 1, column=1, value="Los indicadores son agregados: no incluyen datos personales.").font = Font(
        italic=True, size=9, color="78716C"
    )

    sheets: list[tuple[str, list[str], list[list[Any]], list[int]]] = [
        ("Por día", DAY_HEADERS, _day_rows(summary), [12, 26, 11, 12, 11, 12, 11]),
        (
            "Por médico",
            ["Médico", "Atendidos", "Atención promedio (min)"],
            [[d["doctor"], d["attended"], d["avg_service_minutes"]] for d in summary["by_doctor"]],
            [40, 12, 22],
        ),
        (
            "Por dependencia",
            ["Dependencia", "Solicitudes"],
            [[d["department"], d["count"]] for d in summary["by_department"]],
            [50, 12],
        ),
        (
            "Por canal",
            ["Canal", "Solicitudes"],
            [[CHANNEL_LABEL.get(d["channel"], d["channel"]), d["count"]] for d in summary["by_channel"]],
            [20, 12],
        ),
        (
            "Por hora",
            ["Hora de registro", "Solicitudes"],
            [[f"{d['hour']:02d}:00 – {d['hour']:02d}:59", d["count"]] for d in summary["by_hour"]],
            [20, 12],
        ),
    ]
    rating = summary.get("rating")
    if rating and rating["count"]:
        sheets.append(
            (
                "Satisfacción",
                ["Indicador", "Valor"],
                _rating_rows(rating),
                [40, 30],
            )
        )
    for name, headers, rows, widths in sheets:
        sheet = wb.create_sheet(name)
        start = _xl_title(sheet, f"{name} — Tópico de Salud", meta)
        _xl_table(sheet, start, headers, rows, widths)
        sheet.freeze_panes = sheet.cell(row=start + 1, column=1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# =============================================================================== PDF
_STYLES = {
    "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=INK),
    "meta": ParagraphStyle("meta", fontName="Helvetica", fontSize=9, leading=12, textColor=MUTED),
    "h2": ParagraphStyle(
        "h2", fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=BRAND, spaceBefore=10, spaceAfter=5
    ),
    "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=8, leading=10, textColor=INK, alignment=TA_LEFT),
    "note": ParagraphStyle("note", fontName="Helvetica-Oblique", fontSize=8, leading=10, textColor=MUTED),
}


def _page_decorator(meta: ReportMeta, title: str) -> Any:
    def draw(canvas: Any, doc: Any) -> None:
        width, height = doc.pagesize
        canvas.saveState()
        canvas.setFillColor(BRAND_DARK)
        canvas.rect(0, height - 16 * mm, width, 16 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(15 * mm, height - 9 * mm, meta.institution.upper())
        canvas.setFont("Helvetica", 8.5)
        canvas.drawString(15 * mm, height - 13 * mm, f"Tópico de Salud · {title}")
        canvas.setFillColor(MUTED)
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(15 * mm, 9 * mm, f"Generado el {meta.generated}")
        canvas.drawRightString(width - 15 * mm, 9 * mm, f"Página {doc.page}")
        canvas.setStrokeColor(LINE)
        canvas.line(15 * mm, 13 * mm, width - 15 * mm, 13 * mm)
        canvas.restoreState()

    return draw


def _pdf_table(
    headers: list[str], rows: list[list[Any]], widths: list[float], numeric_from: int = 99, wrap_from: int = 28
) -> Table:
    def fmt(value: Any) -> Any:
        if isinstance(value, datetime):
            return value.astimezone(INSTITUTION_TZ).strftime("%d/%m %H:%M")
        if isinstance(value, date):
            return value.strftime("%d/%m/%Y")
        if isinstance(value, float):
            return f"{value:.1f}"
        if value is None:
            return "—"
        if isinstance(value, str) and len(value) > wrap_from:
            return Paragraph(value, _STYLES["cell"])
        return value

    data = [headers] + [[fmt(v) for v in r] for r in rows]
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), BRAND),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, BRAND_SOFT]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, LINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if numeric_from < len(headers):
        style.append(("ALIGN", (numeric_from, 0), (-1, -1), "RIGHT"))
    table.setStyle(TableStyle(style))
    return table


def _kpi_grid(totals: dict[str, Any], width: float) -> Table:
    cells = []
    for label, value in _kpis(totals):
        cells.append(
            [
                Paragraph(f'<font size="7.5" color="#57534E">{label.upper()}</font>', _STYLES["cell"]),
                Paragraph(f'<font size="14"><b>{value}</b></font>', _STYLES["cell"]),
            ]
        )
    rows = [cells[i : i + 4] for i in range(0, len(cells), 4)]
    data = [[Table([[c[0]], [c[1]]], colWidths=[width / 4 - 6]) for c in row] for row in rows]
    grid = Table(data, colWidths=[width / 4] * 4)
    grid.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.5, LINE),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, LINE),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFAF9")),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return grid


def _daily_chart(summary: dict[str, Any], width: float) -> Drawing | None:
    per_day: dict[date, int] = {}
    for d in summary["by_day"]:
        per_day[d["service_date"]] = per_day.get(d["service_date"], 0) + d["attended"]
    if not per_day:
        return None
    days = sorted(per_day)
    drawing = Drawing(width, 120)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 28, 22, width - 40, 88
    chart.data = [[per_day[d] for d in days]]
    chart.bars[0].fillColor = colors.HexColor("#93263A")
    chart.bars[0].strokeColor = None
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labels.fontName = "Helvetica"
    chart.valueAxis.labels.fontSize = 7
    chart.valueAxis.strokeColor = LINE
    chart.valueAxis.gridStrokeColor = LINE
    chart.valueAxis.visibleGrid = True
    step = max(1, len(days) // 15)
    chart.categoryAxis.categoryNames = [f"{d:%d/%m}" if i % step == 0 else "" for i, d in enumerate(days)]
    chart.categoryAxis.labels.fontName = "Helvetica"
    chart.categoryAxis.labels.fontSize = 6.5
    chart.categoryAxis.strokeColor = LINE
    chart.barSpacing = 1
    drawing.add(chart)
    return drawing


def summary_pdf(summary: dict[str, Any], meta: ReportMeta) -> bytes:
    out = io.BytesIO()
    title = "Reporte de atenciones"
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=24 * mm,
        bottomMargin=18 * mm,
        title=f"{title} {meta.period}",
        author=meta.institution,
    )
    width = doc.width
    story: list[Any] = [
        Paragraph("Reporte de atenciones del Tópico de Salud", _STYLES["title"]),
        Spacer(1, 3),
        Paragraph(f"Periodo: <b>{meta.period}</b> &nbsp;·&nbsp; Sede(s): <b>{meta.sites}</b>", _STYLES["meta"]),
        Spacer(1, 8),
        _kpi_grid(summary["totals"], width),
        Spacer(1, 2),
        Paragraph("Indicadores agregados: no incluyen datos personales ni clínicos.", _STYLES["note"]),
    ]
    chart = _daily_chart(summary, width)
    if chart is not None:
        story += [Paragraph("Atendidos por día", _STYLES["h2"]), chart]

    def section(name: str, table: Table) -> None:
        story.append(KeepTogether([Paragraph(name, _STYLES["h2"]), table]))

    rating = summary.get("rating")
    if rating and rating["count"]:
        section(
            "Satisfacción del servicio (calificación anónima)",
            _pdf_table(["Indicador", "Valor"], _rating_rows(rating), [width * 0.6, width * 0.4], numeric_from=1),
        )
    if summary["by_doctor"]:
        section(
            "Atenciones por médico",
            _pdf_table(
                ["Médico", "Atendidos", "Atención promedio (min)"],
                [[d["doctor"], d["attended"], d["avg_service_minutes"]] for d in summary["by_doctor"]],
                [width * 0.56, width * 0.18, width * 0.26],
                numeric_from=1,
            ),
        )
    half = (width - 8 * mm) / 2
    side_by_side = Table(
        [
            [
                [
                    Paragraph("Por canal de registro", _STYLES["h2"]),
                    _pdf_table(
                        ["Canal", "Solicitudes"],
                        [[CHANNEL_LABEL.get(d["channel"], d["channel"]), d["count"]] for d in summary["by_channel"]]
                        or [["—", 0]],
                        [half * 0.65, half * 0.35],
                        numeric_from=1,
                    ),
                ],
                [
                    Paragraph("Por hora de registro", _STYLES["h2"]),
                    _pdf_table(
                        ["Hora", "Solicitudes"],
                        [[f"{d['hour']:02d}:00 – {d['hour']:02d}:59", d["count"]] for d in summary["by_hour"]]
                        or [["—", 0]],
                        [half * 0.65, half * 0.35],
                        numeric_from=1,
                    ),
                ],
            ]
        ],
        colWidths=[half + 4 * mm, half + 4 * mm],
    )
    side_by_side.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story.append(side_by_side)
    if summary["by_department"]:
        section(
            "Solicitudes por dependencia (principales)",
            _pdf_table(
                ["Dependencia", "Solicitudes"],
                [[d["department"], d["count"]] for d in summary["by_department"]],
                [width * 0.8, width * 0.2],
                numeric_from=1,
            ),
        )
    story.append(Paragraph("Detalle por día y sede", _STYLES["h2"]))
    story.append(
        _pdf_table(
            DAY_HEADERS,
            _day_rows(summary) or [["—", "Sin atenciones en el periodo", 0, 0, 0, 0, 0]],
            [width * w for w in (0.13, 0.27, 0.12, 0.12, 0.12, 0.12, 0.12)],
            numeric_from=2,
        )
    )
    deco = _page_decorator(meta, title)
    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    return out.getvalue()


def nominal_pdf(rows: list[dict[str, Any]], meta: ReportMeta) -> bytes:
    out = io.BytesIO()
    title = "Listado nominal de atenciones"
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=24 * mm,
        bottomMargin=18 * mm,
        title=f"{title} {meta.period}",
        author=meta.institution,
    )
    width = doc.width
    headers = [
        "Fecha",
        "Sede",
        "Turno",
        "DNI",
        "Trabajador",
        "Dependencia",
        "Estado",
        "Registro",
        "Inicio",
        "Fin",
        "Médico",
    ]
    data = [
        [
            r["fecha"],
            r["sede"],
            r["turno"],
            r["dni"],
            f"{r['apellidos']}, {r['nombres']}",
            r["dependencia"],
            STATUS_LABEL.get(r["estado"], r["estado"]),
            r["registrado"],
            r["inicio"],
            r["fin"],
            r["medico"],
        ]
        for r in rows
    ]
    ratios = (0.07, 0.09, 0.05, 0.065, 0.17, 0.17, 0.075, 0.07, 0.07, 0.07, 0.11)
    story: list[Any] = [
        Paragraph(title, _STYLES["title"]),
        Spacer(1, 3),
        Paragraph(
            f"Periodo: <b>{meta.period}</b> &nbsp;·&nbsp; Sede(s): <b>{meta.sites}</b> &nbsp;·&nbsp; "
            f"{len(rows)} registro(s)",
            _STYLES["meta"],
        ),
        Paragraph(
            "Documento con datos personales: uso exclusivo para la gestión del servicio. La descarga quedó registrada.",
            _STYLES["note"],
        ),
        Spacer(1, 8),
        _pdf_table(headers, data or [["—"] * len(headers)], [width * r for r in ratios], wrap_from=14),
    ]
    deco = _page_decorator(meta, title)
    doc.build(story, onFirstPage=deco, onLaterPages=deco)
    return out.getvalue()
