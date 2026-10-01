"""Documento Excel de los datos de DEMOSTRACIÓN: accesos, trabajadores y atenciones.

Lo genera seed_demo al terminar. Contiene contraseñas de prueba: no se versiona
(DATOS-DE-PRUEBA.xlsx está en .gitignore) y se descarta al pasar a producción real.
"""

import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.clock import INSTITUTION_TZ
from app.modules.appointments.models import Appointment, AppointmentStatus
from app.modules.sites.models import Site
from app.modules.workers.models import Worker, WorkerCoverage

BRAND = "7A1E2C"
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
HEADER_FONT = Font(bold=True, color="FFFFFF")
NOTE_FILL = PatternFill("solid", fgColor="FEF3C7")
ROLE_LABEL = {
    "ADMIN": "Administrador(a)",
    "SUPERVISOR": "Supervisor(a)",
    "OPERATOR": "Encargada del tópico",
    "AUDITOR": "Auditor(a)",
}
ROLE_DOES = {
    "ADMIN": "Usuarios, sedes, parámetros, catálogos, importaciones, auditoría",
    "SUPERVISOR": "Reportes y exportaciones, ajuste de capacidad, opera la cola",
    "OPERATOR": "Mesa de atención: registrar, llamar, atender, finalizar",
    "AUDITOR": "Solo lectura: auditoría y reportes",
}
CHANNEL = {"PHONE": "Teléfono", "WALK_IN": "Presencial"}


def _time(value: dt.datetime | None) -> str:
    return value.astimezone(INSTITUTION_TZ).strftime("%H:%M") if value else ""


def _table(ws: Worksheet, headers: list[str], rows: list[list[object]], widths: list[int], start_row: int = 1) -> None:
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for c, value in enumerate(row, 1):
            ws.cell(row=r, column=c, value=value)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def write_report(
    session_factory: object,
    path: str,
    *,
    users: list[tuple[str, str, str, tuple[str, ...]]],
    password: str,
    base_url: str,
    test_dni: str,
    test_email: str | None,
    uncovered_dni: str,
) -> Path:
    wb = Workbook()
    with session_factory() as s:  # type: ignore[operator]
        db: Session = s
        sites = {site.code: site.name for site in db.scalars(select(Site))}
        covered_ids = set(db.scalars(select(WorkerCoverage.worker_id)))
        workers = list(db.scalars(select(Worker).order_by(Worker.paternal_surname, Worker.first_names)))
        labels = {st.code: st.label for st in db.scalars(select(AppointmentStatus))}
        appointments = list(
            db.scalars(
                select(Appointment).order_by(
                    Appointment.service_date.desc(), Appointment.site_id, Appointment.ticket_number
                )
            ).unique()
        )

        # 1. Accesos
        ws = wb.active
        assert ws is not None
        ws.title = "Accesos"
        ws["A1"] = "Sistema del Tópico de Salud — DATOS DE PRUEBA"
        ws["A1"].font = Font(bold=True, size=14, color=BRAND)
        ws["A2"] = f"Dirección del sistema: {base_url or 'http://<servidor>:42000'}"
        ws["A3"] = f"Pantalla de turnos (TV): {base_url}/pantalla/alz  ·  {base_url}/pantalla/bar"
        ws["A4"] = f"Consulta del trabajador (celular): {base_url}/consulta"
        ws["A5"] = "Los usuarios ingresan con su DNI. Datos ficticios: se eliminan antes del uso real."
        ws["A5"].fill = NOTE_FILL
        site_names = lambda codes: ", ".join(sites.get(c, c) for c in codes)  # noqa: E731
        rows: list[list[object]] = [
            [dni, name, ROLE_LABEL[role], site_names(codes), password, ROLE_DOES[role]]
            for dni, name, role, codes in users
        ]
        rows.append(
            [
                "admin",
                "Administrador del Sistema (real)",
                ROLE_LABEL["ADMIN"],
                "Todas",
                "(la que usted definió en el primer ingreso)",
                "Administrador definitivo de la institución",
            ]
        )
        _table(
            ws,
            ["Usuario (DNI)", "Nombre", "Rol", "Sede(s)", "Contraseña", "Qué puede hacer"],
            rows,
            [16, 34, 22, 40, 26, 58],
            start_row=7,
        )

        # 2. Trabajadores
        ws = wb.create_sheet("Trabajadores")
        rows = []
        for w in workers:
            note = ""
            if w.document_number == test_dni:
                note = f"DNI DE PRUEBA: recibe los correos{f' en {test_email}' if test_email else ''}"
            elif w.document_number == uncovered_dni or w.id not in covered_ids:
                note = "Sin cobertura EPS: el sistema rechaza su registro"
            rows.append(
                [
                    w.document_number,
                    f"{w.paternal_surname} {w.maternal_surname or ''}".strip(),
                    w.first_names,
                    w.department.name if w.department else "",
                    w.institutional_email or "",
                    "Sí" if w.id in covered_ids else "No",
                    note,
                ]
            )
        _table(
            ws,
            ["DNI", "Apellidos", "Nombres", "Dependencia", "Correo", "Cobertura EPS", "Observación"],
            rows,
            [12, 26, 20, 40, 32, 14, 50],
        )
        for r, row in enumerate(rows, 2):
            if row[-1]:
                for c in range(1, 8):
                    ws.cell(row=r, column=c).fill = NOTE_FILL

        # 3. Atenciones
        ws = wb.create_sheet("Atenciones")
        rows = [
            [
                a.service_date.strftime("%d/%m/%Y"),
                sites.get(a.site.code, a.site.code),
                a.ticket_code,
                a.worker.document_number,
                f"{a.worker.paternal_surname} {a.worker.maternal_surname or ''}, {a.worker.first_names}".replace(
                    " ,", ","
                ),
                labels.get(a.status, a.status),
                CHANNEL.get(a.channel, a.channel),
                _time(a.registered_at),
                _time(a.called_at),
                _time(a.started_at),
                _time(a.finished_at),
                a.close_reason.label if a.close_reason else "",
            ]
            for a in appointments
        ]
        _table(
            ws,
            [
                "Fecha",
                "Sede",
                "Turno",
                "DNI",
                "Trabajador",
                "Estado",
                "Canal",
                "Registro",
                "Llamado",
                "Inicio",
                "Fin",
                "Motivo",
            ],
            rows,
            [12, 30, 9, 11, 40, 16, 12, 10, 10, 10, 10, 34],
        )

    # 4. Guía de prueba
    ws = wb.create_sheet("Cómo probar")
    steps = [
        f"1. Ingrese con el usuario {users[2][0]} (Encargada Alzamora) y la contraseña {password}.",
        f"2. En otro equipo o ventana abra {base_url}/pantalla/alz y pulse 'Iniciar pantalla'.",
        "3. En la Mesa de atención pulse 'Llamar siguiente' (F4): la pantalla suena y muestra el turno.",
        "4. Pulse 'Iniciar' y luego 'Finalizar' en el turno llamado.",
        f"5. Registre el DNI {test_dni}: turno nuevo" + (f" y correos en {test_email}." if test_email else "."),
        f"6. Intente registrar el DNI {uncovered_dni}: debe rechazarlo (sin cobertura EPS).",
        f"7. En el celular abra {base_url}/consulta con un DNI y turno de la hoja 'Atenciones' de hoy.",
        f"8. Ingrese con {users[1][0]} (Supervisora): Reportes de los últimos días y exportación a Excel.",
        f"9. Ingrese con {users[4][0]} (Auditor): Auditoría con todas las acciones registradas.",
    ]
    ws.column_dimensions["A"].width = 110
    for r, text in enumerate(steps, 1):
        ws.cell(row=r, column=1, value=text)

    target = Path(path).resolve()
    wb.save(target)
    return target
